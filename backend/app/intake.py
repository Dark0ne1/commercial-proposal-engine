"""Untrusted documents are stored verbatim; only text is extracted in the worker."""

import hashlib
import os
import tempfile
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from docx import Document
from docx.table import Table
from openpyxl import load_workbook
from pypdf import PdfReader

from app.application import Problem, create_inquiry
from app.config import settings
from app.domain import InquiryInput


def store_upload(
    db, customer_id: int, filename: str, content: bytes, key: str, actor: str, text=""
):
    name = Path(filename.replace("\\", "/")).name[:200]
    extension = Path(name).suffix.lower()
    if extension not in {".xlsx", ".docx", ".pdf"}:
        raise Problem(422, "Поддерживаются только XLSX, DOCX и текстовые PDF")
    if not content or len(content) > settings().upload_max_bytes:
        raise Problem(413, "Пустой файл или превышен лимит загрузки")
    if extension == ".pdf" and not content.startswith(b"%PDF-"):
        raise Problem(422, "Содержимое не соответствует PDF")
    if extension != ".pdf":
        validate_archive(content, extension)
    digest = hashlib.sha256(content).hexdigest()
    directory = settings().storage_dir / "incoming"
    directory.mkdir(parents=True, exist_ok=True)
    destination = directory / f"{digest}{extension}"
    if not destination.exists():
        with tempfile.NamedTemporaryFile(dir=directory, delete=False) as temporary:
            temporary.write(content)
            temporary_path = Path(temporary.name)
        try:
            os.replace(temporary_path, destination)
        finally:
            temporary_path.unlink(missing_ok=True)
    source = {
        "filename": name,
        "sha256": digest,
        "extension": extension,
        "size": len(content),
        "path": f"incoming/{digest}{extension}",
    }
    return create_inquiry(
        db,
        InquiryInput(customer_id=customer_id, text=text or "Заявка из загруженного документа"),
        key,
        actor,
        source,
    )


def validate_archive(content, extension):
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 1000 or sum(e.file_size for e in entries) > 30 * 1024 * 1024:
                raise Problem(422, "Слишком большой распакованный документ")
            if any(
                e.flag_bits & 1
                or "vbaproject" in e.filename.lower()
                or e.file_size > max(e.compress_size * 200, 1024 * 1024)
                for e in entries
            ):
                raise Problem(422, "Зашифрованный, макросодержащий или чрезмерно сжатый документ")
            required = "word/document.xml" if extension == ".docx" else "xl/workbook.xml"
            if required not in archive.namelist():
                raise Problem(422, "Содержимое не соответствует расширению документа")
            for entry in entries:
                if entry.filename.endswith(".xml") and b"<!DOCTYPE" in archive.read(entry):
                    raise Problem(422, "DTD в документах запрещены")
    except BadZipFile:
        raise Problem(422, "Повреждённый ZIP-контейнер документа") from None


def extract_document(source: dict) -> str:
    path = settings().storage_dir / source["path"]
    expected = settings().storage_dir / "incoming" / f"{source['sha256']}{source['extension']}"
    if path.resolve() != expected.resolve() or not path.resolve().is_relative_to(
        settings().storage_dir.resolve()
    ):
        raise Problem(422, "Неверный путь исходного документа")
    content = path.read_bytes()
    if (
        hashlib.sha256(content).hexdigest() != source["sha256"]
        or len(content) > settings().upload_max_bytes
    ):
        raise Problem(422, "Исходный документ изменён или превышает лимит")
    pieces = []

    def add(value):
        pieces.append(value)
        if sum(len(p) + 1 for p in pieces) > 20000:
            raise Problem(422, "Документ содержит больше 20000 символов; разделите заявку")

    try:
        if source["extension"] in {".docx", ".xlsx"}:
            validate_archive(content, source["extension"])
        if source["extension"] == ".docx":
            for block in Document(BytesIO(content)).iter_inner_content():
                if isinstance(block, Table):
                    for row in block.rows:
                        add("\t".join(cell.text for cell in row.cells))
                else:
                    add(block.text)
        elif source["extension"] == ".xlsx":
            book = load_workbook(
                BytesIO(content), read_only=True, data_only=False, keep_links=False
            )
            try:
                if len(book.worksheets) > 10:
                    raise Problem(422, "Разрешено не более 10 листов")
                cells = 0
                for sheet in book:
                    # Dimensions are untrusted metadata; count the actual streamed cells.
                    sheet.reset_dimensions()
                    for row_number, row in enumerate(sheet.iter_rows(), 1):
                        if row_number > 1000 or len(row) > 50:
                            raise Problem(422, "Превышены размеры листа (1000 × 50)")
                        cells += len(row)
                        if cells > 5000:
                            raise Problem(422, "Разрешено не более 5000 ячеек")
                        if any(cell.data_type == "f" for cell in row):
                            raise Problem(422, "Замените формулы значениями перед загрузкой")
                        add(
                            "\t".join(
                                str(cell.value) if cell.value is not None else "" for cell in row
                            )
                        )
            finally:
                book.close()
        else:
            reader = PdfReader(BytesIO(content))
            if reader.is_encrypted or len(reader.pages) > 50:
                raise Problem(422, "Зашифрованный PDF или больше 50 страниц")
            for page in reader.pages:
                add(page.extract_text() or "")
    except Problem:
        raise
    except Exception:
        raise Problem(422, "Документ повреждён или не поддерживается") from None
    text = "\n".join(pieces).strip()
    if not text:
        raise Problem(
            422,
            "Нет извлекаемого текста. Для скана нужен OCR; загрузите текстовый PDF или вставьте текст",
        )
    return text
