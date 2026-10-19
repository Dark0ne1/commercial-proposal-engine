import logging
import os
import subprocess
import tempfile
import time
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from app.application import Problem, event, inquiry, version_at
from app.config import settings

logger = logging.getLogger(__name__)


def write_docx(snapshot: dict, path: Path):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(1.6)
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = "DejaVu Sans"
        style.font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Normal"].font.size = Pt(9)
    for border in doc.styles["Title"].element.xpath("./w:pPr/w:pBdr"):
        border.getparent().remove(border)
    doc.add_paragraph(snapshot["supplier"]["name"], "Heading 2")
    doc.add_paragraph(snapshot["supplier"]["details"])
    doc.add_heading("Коммерческое предложение", 0)
    doc.add_paragraph(f"{snapshot['number']} от {snapshot['date']}")
    doc.add_paragraph(f"Клиент: {snapshot['customer']['name']}\n{snapshot['customer']['details']}")
    table = doc.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    table.autofit = False
    widths = [1.0, 6.0, 2.0, 2.5, 2.0, 3.0]
    for column, width in zip(table.columns, widths, strict=True):
        column.width = Cm(width)
    headers = ["№", "Наименование / артикул", "Кол-во", "Цена, ₽", "Скидка, ₽", "Сумма, ₽"]
    for cell, label, width in zip(table.rows[0].cells, headers, widths, strict=True):
        cell.text = label
        cell.width = Cm(width)
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "E8EDF1")
        cell._tc.get_or_add_tcPr().append(shading)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for index, line in enumerate(snapshot["lines"], 1):
        row = table.add_row()
        values = [
            str(index),
            f"{line['name']}\n{line['sku']}\n{line['availability']}"
            + (
                f"\nПроизводство: {line['production_days']} дн.; доставка: {line['delivery_days']} дн."
                if "production_days" in line
                else ""
            ),
            f"{line['quantity']} {line['unit']}",
            line["unit_price"],
            line["discount_amount"],
            line["amount"],
        ]
        for cell, value, width in zip(row.cells, values, widths, strict=True):
            cell.text = value
            cell.width = Cm(width)
        cant_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant_split)
    summary_start = len(doc.paragraphs)
    for label, value in [
        ("До скидки", snapshot["subtotal"]),
        (f"Скидка {snapshot['discount']}%", snapshot["discount_amount"]),
        ("Без НДС", snapshot["net"]),
        (f"НДС {snapshot['vat_rate']}%", snapshot["vat"]),
        ("Итого к оплате", snapshot["total"]),
    ]:
        p = doc.add_paragraph(f"{label}: {value} RUB")
        if label == "Итого к оплате":
            p.runs[0].bold = True
    doc.add_paragraph(f"Условия оплаты: {snapshot['payment_terms']}")
    doc.add_paragraph(f"Предложение действительно до {snapshot['valid_until']}")
    if snapshot["comments"]:
        doc.add_paragraph(f"Комментарий менеджера: {snapshot['comments']}")
    doc.add_paragraph(snapshot["notice"])
    for paragraph in doc.paragraphs[summary_start:-1]:
        paragraph.paragraph_format.keep_with_next = True
    doc.save(path)


def render(snapshot: dict, target: Path):
    config = settings()
    target.mkdir(parents=True, exist_ok=True)
    # Each process uses its own LibreOffice profile and atomic output replacement.
    with (
        tempfile.TemporaryDirectory(dir=target) as temporary,
        tempfile.TemporaryDirectory(prefix="lo-") as profile,
    ):
        work = Path(temporary).resolve()
        source = work / "proposal.docx"
        write_docx(snapshot, source)
        result = subprocess.run(
            [
                config.libreoffice_bin,
                f"-env:UserInstallation={Path(profile).resolve().as_uri()}",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(work),
                str(source),
            ],
            capture_output=True,
            timeout=config.render_timeout,
            check=False,
        )
        pdf = work / "proposal.pdf"
        if result.returncode or not pdf.exists() or not pdf.read_bytes().startswith(b"%PDF"):
            logger.error(
                "LibreOffice exit=%s stdout=%r stderr=%r",
                result.returncode,
                result.stdout,
                result.stderr,
            )
            raise RuntimeError("LibreOffice не создал PDF")
        os.replace(source, target / "proposal.docx")
        os.replace(pdf, target / "proposal.pdf")


def artifact_directory(version_id):
    return settings().storage_dir / version_id


def generate_documents(db, identifier, actor):
    item = inquiry(db, identifier, lock=True)
    if not item.confirmed_revision:
        raise Problem(409, "Сначала подтвердите КП")
    version = version_at(db, identifier, item.confirmed_revision)
    target = artifact_directory(version.id)
    if not all((target / f"proposal.{ext}").is_file() for ext in ("docx", "pdf")):
        start = time.monotonic()
        try:
            render(version.snapshot, target)
        except (OSError, RuntimeError, subprocess.TimeoutExpired) as error:
            db.rollback()
            logger.exception("Document rendering failed inquiry=%s", identifier)
            event(db, item, actor, "documents.failed", {"error": type(error).__name__})
            db.commit()
            raise Problem(
                503, "Не удалось создать документы. Подтверждение сохранено; повторите генерацию"
            ) from error
        event(
            db,
            item,
            actor,
            "documents.generated",
            {"revision": version.revision, "duration_ms": round((time.monotonic() - start) * 1000)},
        )
        db.commit()
    return {
        "docx": f"/api/inquiries/{identifier}/files/docx",
        "pdf": f"/api/inquiries/{identifier}/files/pdf",
    }
