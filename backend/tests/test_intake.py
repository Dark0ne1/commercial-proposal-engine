from io import BytesIO
from zipfile import ZipFile

import pytest
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter
from sqlalchemy import func, select

from app.application import Problem
from app.config import settings
from app.intake import extract_document
from app.llm import MockProvider, ProviderError
from app.models import Inquiry
from app.processing import ProcessInput, create_run, perform


def document_bytes(kind):
    output = BytesIO()
    if kind == "docx":
        document = Document()
        document.add_paragraph("SKU-0001, 2 шт.")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "Условия", "100% предоплата"
        document.save(output)
    else:
        book = Workbook()
        book.active.append(["SKU-0001", 2, "шт"])
        book.save(output)
    return output.getvalue()


@pytest.mark.parametrize("kind", ["docx", "xlsx"])
def test_import_to_agent_draft(manager, db, kind):
    content = document_bytes(kind)
    response = manager.post(
        "/api/inquiries/import",
        data={"customer_id": "1"},
        files={"file": (f"../../source.{kind}", content)},
        headers={"Idempotency-Key": "import-test-key"},
    )
    assert response.status_code == 201, response.text
    data = response.json()
    assert data["source_info"]["filename"] == f"source.{kind}"
    assert "path" not in data["source_info"]
    repeat = manager.post(
        "/api/inquiries/import",
        data={"customer_id": "1"},
        files={"file": (f"../../source.{kind}", content)},
        headers={"Idempotency-Key": "import-test-key"},
    )
    assert repeat.json()["id"] == data["id"]
    item = db.get(Inquiry, data["id"])
    run = create_run(db, item.id, ProcessInput(), "import-run-key", "manager")
    assert perform(db, run.id, MockProvider())["status"] == "completed"
    db.refresh(item)
    assert "SKU-0001" in item.extracted_text
    assert item.confirmed_revision is None
    assert db.scalar(select(func.count()).select_from(Inquiry)) == 1


def test_reject_wrong_file_type_and_signature(manager):
    for name, content in [
        ("payload.exe", b"MZ"),
        ("payload.pdf", b"not a pdf"),
        ("payload.xlsx", b"PK invalid"),
    ]:
        response = manager.post(
            "/api/inquiries/import",
            data={"customer_id": "1"},
            files={"file": (name, content)},
            headers={"Idempotency-Key": "invalid-file-key"},
        )
        assert response.status_code == 422


def test_upload_limit(manager, monkeypatch):
    monkeypatch.setenv("UPLOAD_MAX_BYTES", "1024")
    settings.cache_clear()
    response = manager.post(
        "/api/inquiries/import",
        data={"customer_id": "1"},
        files={"file": ("source.pdf", b"%PDF-" + b"x" * 1024)},
        headers={"Idempotency-Key": "large-file-key"},
    )
    assert response.status_code == 413


def test_formula_and_scan_need_manual_correction(manager, db):
    book = Workbook()
    book.active.append(["SKU-0001", "=1+1", "шт"])
    content = BytesIO()
    book.save(content)
    scan = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    writer.write(scan)
    for index, (name, value) in enumerate(
        [("formula.xlsx", content.getvalue()), ("scan.pdf", scan.getvalue())]
    ):
        response = manager.post(
            "/api/inquiries/import",
            data={"customer_id": "1"},
            files={"file": (name, value)},
            headers={"Idempotency-Key": f"unsupported-content-{index}"},
        )
        assert response.status_code == 201
        run = create_run(
            db, response.json()["id"], ProcessInput(), f"unsupported-run-{index}", "manager"
        )
        with pytest.raises(ProviderError, match="формул|OCR"):
            perform(db, run.id, MockProvider())
        assert run.status == "failed"


def test_archive_bomb_and_dtd_are_rejected(manager):
    cases = []
    bomb = BytesIO()
    with ZipFile(bomb, "w", compression=8) as archive:
        archive.writestr("word/document.xml", b"x" * (2 * 1024 * 1024))
    cases.append(bomb.getvalue())
    dtd = BytesIO()
    with ZipFile(dtd, "w") as archive:
        archive.writestr("word/document.xml", '<!DOCTYPE x [<!ENTITY a "x">]><x>&a;</x>')
    cases.append(dtd.getvalue())
    for index, content in enumerate(cases):
        response = manager.post(
            "/api/inquiries/import",
            data={"customer_id": "1"},
            files={"file": ("source.docx", content)},
            headers={"Idempotency-Key": f"bad-archive-key-{index}"},
        )
        assert response.status_code == 422


def test_source_digest_prevents_modified_file(manager, db):
    response = manager.post(
        "/api/inquiries/import",
        data={"customer_id": "1"},
        files={"file": ("source.docx", document_bytes("docx"))},
        headers={"Idempotency-Key": "source-digest-key"},
    )
    item = db.get(Inquiry, response.json()["id"])
    path = settings().storage_dir / item.source_info["path"]
    path.write_bytes(b"tampered")
    with pytest.raises(Problem):
        extract_document(item.source_info)


def test_forged_xlsx_dimensions_cannot_hide_positions(manager, db):
    book = Workbook()
    book.active.append(["SKU-0001", 2, "шт"])
    book.active.append(["SKU-0002", 1, "шт"])
    original = BytesIO()
    book.save(original)
    forged = BytesIO()
    with ZipFile(original) as source, ZipFile(forged, "w") as target:
        for entry in source.infolist():
            content = source.read(entry)
            if entry.filename == "xl/worksheets/sheet1.xml":
                content = content.replace(b'ref="A1:C2"', b'ref="A1:A1"')
            target.writestr(entry, content)
    response = manager.post(
        "/api/inquiries/import",
        data={"customer_id": "1"},
        files={"file": ("request.xlsx", forged.getvalue())},
        headers={"Idempotency-Key": "forged-dimension-key"},
    )
    item = db.get(Inquiry, response.json()["id"])
    assert "SKU-0002" in extract_document(item.source_info)


def test_text_pdf_to_draft(manager, db, tmp_path):
    import shutil
    from decimal import Decimal

    from app.documents import render
    from app.domain import calculate

    if not shutil.which(settings().libreoffice_bin):
        pytest.skip("LibreOffice required for PDF fixture")
    target = tmp_path / "pdf-source"
    snapshot = {
        **calculate([], Decimal(0), Decimal(20)),
        "supplier": {"name": "Поставщик", "details": "Реквизиты"},
        "customer": {"name": "Клиент", "details": "Тестовые реквизиты"},
        "number": "SOURCE-1",
        "date": "2026-10-19",
        "valid_until": "2026-11-02",
        "payment_terms": "Аванс",
        "comments": "SKU-0001, 2 шт.",
        "notice": "Тестовый входящий документ",
    }
    render(snapshot, target)
    response = manager.post(
        "/api/inquiries/import",
        data={"customer_id": "1"},
        files={"file": ("request.pdf", (target / "proposal.pdf").read_bytes())},
        headers={"Idempotency-Key": "text-pdf-key"},
    )
    assert response.status_code == 201
    run = create_run(db, response.json()["id"], ProcessInput(), "pdf-process-key", "manager")
    assert perform(db, run.id, MockProvider())["status"] == "completed"
